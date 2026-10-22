"""Export the logbook to spreadsheets.

* xrplarb_export.xlsx - one workbook with tabs: Summary, Tax (one row per
  swap, Koinly "universal" layout which most crypto tax tools import),
  Trades, Missed gaps, Fees, Transfers.
* one CSV per tab as well, for tools that prefer CSV.

Note for your accountant: every loop is two swaps in your wallet
(XRP -> token, token -> XRP). The intermediate hops happen inside other
people's offers/pools and never touch your wallet, so they are not your
disposals. Network fees are listed separately.
"""

from __future__ import annotations

import csv
from decimal import Decimal
from pathlib import Path
from typing import Dict, List, Optional

from .amounts import Asset
from .db import DB


def _asset_label(key: str) -> str:
    if not key:
        return ""
    if key == "XRP":
        return "XRP"
    a = Asset.from_key(key)
    return f"{a.symbol} ({a.issuer})"


def _where(network: Optional[str], col: str = "network") -> tuple[str, list]:
    return (f"WHERE {col}=?", [network]) if network else ("", [])


def tax_rows(db: DB, network: Optional[str]) -> List[Dict[str, str]]:
    w, args = _where(network)
    rows = []
    for r in db.query(f"SELECT * FROM legs {w} ORDER BY ts", args):
        date = r["ts"][:19].replace("T", " ")
        if r["label"] == "network fee":
            rows.append({
                "Date": date, "Sent Amount": "", "Sent Currency": "", "Received Amount": "",
                "Received Currency": "", "Fee Amount": r["fee_xrp"], "Fee Currency": "XRP",
                "Net Worth Amount": "", "Net Worth Currency": "", "Label": "cost",
                "Description": f"{r['bot']} network fee", "TxHash": r["tx_hash"],
            })
            continue
        rows.append({
            "Date": date,
            "Sent Amount": r["sent_amount"] if r["sent_asset"] else "",
            "Sent Currency": _asset_label(r["sent_asset"]),
            "Received Amount": r["received_amount"] if r["received_asset"] else "",
            "Received Currency": _asset_label(r["received_asset"]),
            "Fee Amount": "", "Fee Currency": "",
            "Net Worth Amount": "", "Net Worth Currency": "",
            "Label": "swap" if r["label"] in ("trade", "liquidation", "recovery", "dust-sweep") else r["label"],
            "Description": f"{r['bot']} {r['label']}",
            "TxHash": r["tx_hash"],
        })
    return rows


def summary_rows(db: DB, network: Optional[str]) -> List[Dict[str, str]]:
    w, args = _where(network)
    out = []
    bots = [r["bot"] for r in db.query(f"SELECT DISTINCT bot FROM trades {w} ORDER BY bot", args)]
    for b in bots:
        q = f"SELECT status, profit_xrp, network_fee_xrp FROM trades {w + (' AND' if w else 'WHERE')} bot=?"
        rs = db.query(q, args + [b])
        wins = sum(1 for r in rs if r["status"] == "win")
        losses = sum(1 for r in rs if r["status"] == "loss")
        failed = sum(1 for r in rs if r["status"] == "failed")
        attempts = wins + losses + failed
        pnl = sum((Decimal(r["profit_xrp"] or "0") for r in rs), Decimal(0))
        fees = sum((Decimal(r["network_fee_xrp"] or "0") for r in rs), Decimal(0))
        out.append({
            "Bot": b, "Attempts": attempts, "Wins": wins, "Losses": losses, "Cancelled (fee only)": failed,
            "Win rate": f"{(wins / attempts * 100):.1f}%" if attempts else "",
            "Realised P&L (XRP)": str(pnl), "Network fees (XRP)": str(fees),
        })
    return out


def export_all(db: DB, out_dir: Path, network: Optional[str] = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    out_dir.mkdir(parents=True, exist_ok=True)
    w, args = _where(network)
    sheets: Dict[str, List[Dict]] = {
        "Summary": summary_rows(db, network),
        "Tax": tax_rows(db, network),
        "Trades": [dict(r) for r in db.query(f"SELECT * FROM trades {w} ORDER BY ts", args)],
        "Missed gaps": [
            dict(r) for r in db.query(
                f"SELECT * FROM opportunities {w + (' AND' if w else 'WHERE')} action != 'executed' ORDER BY ts", args
            )
        ],
        "Fees": [dict(r) for r in db.query(f"SELECT ts, bot, tx_hash, fee_xrp FROM legs {w + (' AND' if w else 'WHERE')} label='network fee' ORDER BY ts", args)],
        "Transfers": [dict(r) for r in db.query(f"SELECT * FROM transfers {w} ORDER BY ts", args)],
    }
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        if not rows:
            ws.append(["(no rows yet)"])
        else:
            headers = list(rows[0].keys())
            ws.append(headers)
            for c in ws[1]:
                c.font = Font(bold=True)
            for r in rows:
                ws.append([_cell(r.get(h)) for h in headers])
            for i, h in enumerate(headers, 1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(60, max(10, len(h) + 4))
        with open(out_dir / f"{name.lower().replace(' ', '_')}.csv", "w", newline="") as f:
            if rows:
                cw = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
                cw.writeheader()
                cw.writerows(rows)
    path = out_dir / "xrplarb_export.xlsx"
    wb.save(path)
    return path


def _cell(v):
    """Numbers as numbers (so the spreadsheet can sum them), text as text."""
    if isinstance(v, str):
        try:
            if v and v.replace(".", "", 1).replace("-", "", 1).isdigit():
                return float(Decimal(v))
        except Exception:
            pass
    return v
