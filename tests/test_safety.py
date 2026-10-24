import importlib
import logging
import os
import stat
from decimal import Decimal as D

import pytest


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("XRPLARB_SECRETS_DIR", str(tmp_path / "secrets"))
    monkeypatch.setenv("XRPLARB_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("XRPLARB_CONFIG_DIR", str(tmp_path / "config"))
    (tmp_path / "config").mkdir()
    import xrplarb.config as config
    import xrplarb.secrets_store as ss

    importlib.reload(config)
    importlib.reload(ss)
    for mod in ("xrplarb.db", "xrplarb.dashboard", "xrplarb.alerts"):
        importlib.reload(importlib.import_module(mod))
    return tmp_path


SEED = "sEdTM1uX8pu2do5XvTnutH6HsouMaM2"


def test_secrets_file_is_owner_only(env):
    import xrplarb.secrets_store as ss

    ss.add_wallet("testnet", "bot1", "rADDR", SEED)
    p = ss.secrets_path()
    assert stat.S_IMODE(p.stat().st_mode) == 0o600
    assert stat.S_IMODE(p.parent.stat().st_mode) == 0o700
    assert SEED not in ss.describe()
    os.chmod(p, 0o644)
    with pytest.raises(ss.SecretsError):
        ss.load()


def test_redaction_in_logs_and_db(env, caplog):
    from xrplarb.db import DB
    from xrplarb.redact import RedactingFilter, register_secrets, scrub

    register_secrets([SEED])
    assert SEED not in scrub(f"oops {SEED} leaked")
    assert "sEdUnknownSeedXXXXXXXXXXXXXXXXX" not in scrub("seed sEdUnknownSeedXXXXXXXXXXXXXXXXX")
    # tx hashes and addresses survive
    h = "A" * 64
    assert h in scrub(h) and "rhub8VRN55s94qWKDv6jmDy1pUykJzF3wq" in scrub("rhub8VRN55s94qWKDv6jmDy1pUykJzF3wq")
    lg = logging.getLogger("t")
    lg.addFilter(RedactingFilter())
    with caplog.at_level(logging.INFO):
        lg.info("seed is %s", SEED)
    assert SEED not in caplog.text
    db = DB()
    db.event("testnet", "bot1", "INFO", "x", f"seed {SEED}")
    assert SEED not in db.query("SELECT message FROM events")[0]["message"]


def test_encrypted_backup_roundtrip(env):
    import xrplarb.secrets_store as ss

    ss.add_wallet("testnet", "bot1", "rADDR", SEED)
    out = env / "b.enc"
    ss.encrypt_backup("correct horse battery staple", out)
    assert SEED.encode() not in out.read_bytes()
    assert SEED.encode() in ss.decrypt_backup("correct horse battery staple", out)
    with pytest.raises(ss.SecretsError):
        ss.decrypt_backup("wrong", out)


def test_daily_loss_cap_values_and_ignores_transfers(env):
    from xrplarb.db import DB
    from xrplarb.risk import daily_loss, loss_cap_hit

    db = DB()
    db.balance(network="testnet", bot="bot1", xrp=D(100), tokens_value_xrp=D(0), total_value_xrp=D(100), holdings={})
    # 20 XRP swept to the vault is not a loss
    db.transfer(network="testnet", bot="bot1", direction="out", amount_xrp=D(20), tx_hash="", note="sweep")
    loss, frac, start = daily_loss(db, "testnet", "bot1", D(79))
    assert start == D(100) and loss == D(1) and not loss_cap_hit(frac, D("0.10"))
    loss, frac, _ = daily_loss(db, "testnet", "bot1", D("69.9"))
    assert loss_cap_hit(frac, D("0.10"))


def test_liquidation_value_is_not_face_value():
    from xrplarb.amounts import XRP, Asset
    from xrplarb.market import MarketSnapshot, liquidation_value
    from xrplarb.quote import Edge, Offer

    T = Asset("USD", "rI")
    s = MarketSnapshot(ledger_index=1, taken_at=0)
    s.edges[(T.key, XRP.key)] = Edge(T, XRP, offers=[Offer(out_amount=D(2), in_amount=D(1))])  # only 1 USD of bids
    v = liquidation_value(s, {T: D(100)})
    assert v[T] == D(2)  # 100 tokens, but only 2 XRP of buyers
    s2 = MarketSnapshot(ledger_index=1, taken_at=0)
    assert liquidation_value(s2, {T: D(100)})[T] == 0


def test_settings_enforce_owner_rules(env):
    from xrplarb.config import load_settings

    cfg = env / "config" / "settings.yaml"
    cfg.write_text("working_balance_xrp: 500\n")
    with pytest.raises(ValueError):
        load_settings(cfg)
    cfg.write_text("daily_loss_limit_pct: 0.5\n")
    with pytest.raises(ValueError):
        load_settings(cfg)
    cfg.write_text("num_bots: 9\n")
    with pytest.raises(ValueError):
        load_settings(cfg)
    cfg.write_text("max_trade_xrp: 80\nworking_balance_xrp: 20\n")
    s = load_settings(cfg)
    assert s.max_trade_xrp == 20 and not s.allow_mainnet_trading and s.network == "testnet"


def test_dashboard_requires_password_and_csrf(env):
    from fastapi.testclient import TestClient

    import xrplarb.secrets_store as ss
    from xrplarb.dashboard import create_app
    from xrplarb.db import DB

    ss.set_value("dashboard.password_hash", ss.hash_password("a very long password"))
    ss.add_wallet("testnet", "bot1", "rADDR", SEED)
    db = DB()
    db.balance(network="testnet", bot="bot1", xrp=D(100), tokens_value_xrp=D(0), total_value_xrp=D(100), holdings={})
    c = TestClient(create_app(db, "testnet"))
    r = c.get("/", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"
    assert c.post("/kill", data={"csrf": "x", "confirm": "STOP"}, follow_redirects=False).status_code == 303
    c.post("/login", data={"password": "wrong"})
    assert c.get("/", follow_redirects=False).status_code == 303
    c.post("/login", data={"password": "a very long password"})
    r = c.get("/")
    assert r.status_code == 200 and "bot1" in r.text and SEED not in r.text
    assert c.post("/kill", data={"csrf": "forged", "confirm": "STOP"}).status_code == 400
    assert db.get_flag("kill") is None
    import re

    csrf = re.search(r"name=csrf value=(\w+)", r.text).group(1)
    c.post("/kill", data={"csrf": csrf, "confirm": "nope"})
    assert db.get_flag("kill") is None
    c.post("/kill", data={"csrf": csrf, "confirm": "STOP"})
    assert db.get_flag("kill") == "1"
    c.post("/resume", data={"csrf": csrf})
    assert db.get_flag("kill") is None and db.get_flag("resume") == "1"


def test_login_lockout(env):
    from fastapi.testclient import TestClient

    import xrplarb.secrets_store as ss
    from xrplarb.dashboard import create_app
    from xrplarb.db import DB

    ss.set_value("dashboard.password_hash", ss.hash_password("a very long password"))
    c = TestClient(create_app(DB(), "testnet"))
    for _ in range(5):
        c.post("/login", data={"password": "nope"})
    r = c.post("/login", data={"password": "a very long password"})
    assert "Too many attempts" in r.text


def test_export_xlsx_and_tax_rows(env):
    from openpyxl import load_workbook

    from xrplarb.db import DB
    from xrplarb.export import export_all

    db = DB()
    tid = db.trade(network="testnet", bot="bot1", route="XRP>USD.rI>XRP", mode="atomic", tx_hash="H", status="win",
                   profit_xrp=D("0.05"), network_fee_xrp=D("0.00004"), details={})
    db.leg(trade_id=tid, network="testnet", bot="bot1", tx_hash="H1", sent_asset="XRP", sent_amount=D(10),
           received_asset="USD.rI", received_amount=D(5), fee_xrp=D(0), label="trade")
    db.leg(trade_id=tid, network="testnet", bot="bot1", tx_hash="H", sent_asset="", sent_amount=D(0),
           received_asset="", received_amount=D(0), fee_xrp=D("0.00004"), label="network fee")
    p = export_all(db, env / "exp", "testnet")
    wb = load_workbook(p)
    assert set(wb.sheetnames) >= {"Summary", "Tax", "Trades", "Missed gaps", "Fees"}
    tax = list(wb["Tax"].values)
    assert tax[0][:5] == ("Date", "Sent Amount", "Sent Currency", "Received Amount", "Received Currency")
    assert tax[1][1] == 10 and tax[1][2] == "XRP" and tax[2][5] == 0.00004
    assert (env / "exp" / "tax.csv").exists()


def test_issuer_with_signer_list_is_not_called_locked():
    # RLUSD's issuer: master key disabled but controlled by a multi-sig list.
    from xrplarb.amounts import Asset
    from xrplarb.tokens import LSF_DISABLE_MASTER, TokenReport, assess_issuer

    r = TokenReport(asset=Asset("USD", "rI"))
    assess_issuer(r, {"Flags": LSF_DISABLE_MASTER, "signer_lists": [{"SignerQuorum": 2}]})
    assert not r.issuer_locked
    r2 = TokenReport(asset=Asset("USD", "rI"))
    assess_issuer(r2, {"Flags": LSF_DISABLE_MASTER, "signer_lists": []})
    assert r2.issuer_locked


def test_liquidity_ignores_absurdly_priced_orders():
    from xrplarb.amounts import XRP, Asset
    from xrplarb.market import MarketSnapshot
    from xrplarb.quote import Edge, Offer
    from xrplarb.tokens import TokenReport, assess_liquidity

    T = Asset("BIT", "rI")
    s = MarketSnapshot(ledger_index=1, taken_at=0)
    s.edges[(XRP.key, T.key)] = Edge(XRP, T, offers=[
        Offer(out_amount=D(100), in_amount=D(1000)),               # real: 10 XRP each
        Offer(out_amount=D(1), in_amount=D(500_000_000)),          # absurd price
    ])
    s.edges[(T.key, XRP.key)] = Edge(T, XRP, offers=[Offer(out_amount=D(900), in_amount=D(100))])
    r = TokenReport(asset=T)
    assess_liquidity(r, s, D(5000))
    assert r.xrp_liquidity == D(1000)


def test_live_view_needs_login_and_shows_no_secrets(env):
    from fastapi.testclient import TestClient

    import xrplarb.secrets_store as ss
    from xrplarb.dashboard import create_app
    from xrplarb.db import DB

    ss.set_value("dashboard.password_hash", ss.hash_password("a very long password"))
    ss.add_wallet("testnet", "bot1", "rADDR", SEED)
    db = DB()
    db.balance(network="testnet", bot="bot1", xrp=D(100), tokens_value_xrp=D(0), total_value_xrp=D(100), holdings={})
    db.heartbeat("bot1", "running", "")
    c = TestClient(create_app(db, "testnet"))
    assert c.get("/live", follow_redirects=False).status_code == 303
    assert c.get("/live.json").status_code == 401
    c.post("/login", data={"password": "a very long password"})
    assert "PORTAL OPS" in c.get("/live").text
    j = c.get("/live.json")
    assert j.status_code == 200 and SEED not in j.text
    data = j.json()
    assert data["bots"][0]["name"] == "bot1" and data["paper"]["shown"] is False and data["testnet"]["trades"] == 0
